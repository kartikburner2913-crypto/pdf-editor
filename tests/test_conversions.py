import io
import os
import pytest
import pymupdf as fitz
from docx import Document
from pptx import Presentation
import openpyxl
from fastapi.testclient import TestClient

from app.main import app, DOC_STORE, DATA_DIR
from app.pdf_engine import PDFEngine

client = TestClient(app)


def create_sample_pdf(num_pages: int = 2) -> bytes:
    doc = fitz.open()
    for i in range(num_pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 60), f"Sample Document Page {i+1}", fontsize=16)
        page.insert_text((50, 100), "This is a paragraph of text used for format conversion testing.", fontsize=11)
        page.insert_text((50, 130), "Item A: 100 | Item B: 200 | Item C: 300", fontsize=10)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx() -> bytes:
    doc = Document()
    doc.add_heading("Test Document Title", level=1)
    doc.add_paragraph("This is a test paragraph inside a Word document.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Value A"
    table.cell(1, 1).text = "Value B"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(["Category", "Q1", "Q2", "Total"])
    ws.append(["Revenue", 1000, 1500, 2500])
    ws.append(["Expenses", 400, 500, 900])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_sample_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Presentation Title"
    subtitle.text = "Subtitle for PPTX test"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class TestFormatConversions:
    """Test full document conversions between PDF and Office formats (SmallPDF Parity)."""

    def test_pdf_to_word(self):
        pdf_bytes = create_sample_pdf(1)
        docx_bytes = PDFEngine.convert_pdf_to_docx(pdf_bytes)
        assert len(docx_bytes) > 0
        # Valid docx can be opened by python-docx
        doc = Document(io.BytesIO(docx_bytes))
        assert len(doc.paragraphs) > 0 or len(doc.tables) > 0

        # Test API upload endpoint
        res = client.post(
            "/api/convert-pdf-to-word",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")}
        )
        assert res.status_code == 200
        assert "wordprocessingml.document" in res.headers.get("content-type", "")

    def test_pdf_to_excel(self):
        pdf_bytes = create_sample_pdf(2)
        excel_bytes = PDFEngine.convert_pdf_to_excel(pdf_bytes)
        assert len(excel_bytes) > 0
        # Valid xlsx can be opened by openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        assert len(wb.sheetnames) >= 2
        assert "Page 1" in wb.sheetnames

        # Test API upload endpoint
        res = client.post(
            "/api/convert-pdf-to-excel",
            files={"file": ("report.pdf", pdf_bytes, "application/pdf")}
        )
        assert res.status_code == 200
        assert "spreadsheetml.sheet" in res.headers.get("content-type", "")

    def test_pdf_to_pptx(self):
        pdf_bytes = create_sample_pdf(2)
        pptx_bytes = PDFEngine.convert_pdf_to_pptx(pdf_bytes)
        assert len(pptx_bytes) > 0
        # Valid pptx can be opened by python-pptx
        prs = Presentation(io.BytesIO(pptx_bytes))
        assert len(prs.slides) == 2

        # Test API upload endpoint
        res = client.post(
            "/api/convert-pdf-to-pptx",
            files={"file": ("deck.pdf", pdf_bytes, "application/pdf")}
        )
        assert res.status_code == 200
        assert "presentationml.presentation" in res.headers.get("content-type", "")

    def test_docx_to_pdf(self):
        docx_bytes = create_sample_docx()
        pdf_bytes = PDFEngine.convert_docx_to_pdf(docx_bytes)
        assert len(pdf_bytes) > 0
        assert pdf_bytes.startswith(b"%PDF")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        assert len(doc) >= 1
        doc.close()

        # Test API endpoint
        res = client.post(
            "/api/convert-docx-to-pdf",
            files={"file": ("document.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert res.status_code == 200
        assert "application/pdf" in res.headers.get("content-type", "")
        assert res.content.startswith(b"%PDF")

    def test_excel_to_pdf(self):
        xlsx_bytes = create_sample_xlsx()
        pdf_bytes = PDFEngine.convert_excel_to_pdf(xlsx_bytes)
        assert len(pdf_bytes) > 0
        assert pdf_bytes.startswith(b"%PDF")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        assert len(doc) >= 1
        doc.close()

        # Test API endpoint
        res = client.post(
            "/api/convert-excel-to-pdf",
            files={"file": ("sheet.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        assert res.status_code == 200
        assert "application/pdf" in res.headers.get("content-type", "")
        assert res.content.startswith(b"%PDF")

    def test_pptx_to_pdf(self):
        pptx_bytes = create_sample_pptx()
        pdf_bytes = PDFEngine.convert_pptx_to_pdf(pptx_bytes)
        assert len(pdf_bytes) > 0
        assert pdf_bytes.startswith(b"%PDF")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        assert len(doc) >= 1
        doc.close()

        # Test API endpoint
        res = client.post(
            "/api/convert-pptx-to-pdf",
            files={"file": ("slides.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
        )
        assert res.status_code == 200
        assert "application/pdf" in res.headers.get("content-type", "")
        assert res.content.startswith(b"%PDF")

    def test_pdf_to_pdfa(self):
        pdf_bytes = create_sample_pdf(1)
        pdfa_bytes = PDFEngine.convert_pdf_to_pdfa(pdf_bytes, conformance="2b")
        assert len(pdfa_bytes) > 0
        assert pdfa_bytes.startswith(b"%PDF")
        doc = fitz.open(stream=pdfa_bytes, filetype="pdf")
        assert len(doc) == 1
        # Check XML metadata contains pdfaid schema
        xmp = doc.get_xml_metadata()
        assert "pdfaid" in xmp or doc.metadata.get("producer") == "PDF Studio Engine"
        doc.close()

        # Test API endpoint
        res = client.post(
            "/api/convert-pdf-to-pdfa",
            files={"file": ("archive.pdf", pdf_bytes, "application/pdf")},
            data={"level": "2b"}
        )
        assert res.status_code == 200
        assert "application/pdf" in res.headers.get("content-type", "")
        assert res.content.startswith(b"%PDF")

    def test_active_document_conversion_endpoints(self):
        pdf_bytes = create_sample_pdf(1)
        upload_res = client.post(
            "/api/upload",
            files={"file": ("active_doc.pdf", pdf_bytes, "application/pdf")}
        )
        assert upload_res.status_code == 200
        doc_id = upload_res.json()["doc_id"]

        # 1. Convert active doc to word
        word_res = client.get(f"/api/document/{doc_id}/convert-to-word")
        assert word_res.status_code == 200
        assert "wordprocessingml.document" in word_res.headers.get("content-type", "")

        # 2. Convert active doc to excel
        excel_res = client.get(f"/api/document/{doc_id}/convert-to-excel")
        assert excel_res.status_code == 200
        assert "spreadsheetml.sheet" in excel_res.headers.get("content-type", "")

        # 3. Convert active doc to pptx
        pptx_res = client.get(f"/api/document/{doc_id}/convert-to-pptx")
        assert pptx_res.status_code == 200
        assert "presentationml.presentation" in pptx_res.headers.get("content-type", "")

        # 4. Convert active doc to pdfa
        pdfa_res = client.get(f"/api/document/{doc_id}/convert-to-pdfa?level=2b")
        assert pdfa_res.status_code == 200
        assert "application/pdf" in pdfa_res.headers.get("content-type", "")
        assert pdfa_res.content.startswith(b"%PDF")

        # 5. Conversion Feasibility Assessment API
        feasibility_ppt_res = client.get(f"/api/document/{doc_id}/conversion-feasibility?target=pptx")
        assert feasibility_ppt_res.status_code == 200
        ppt_meta = feasibility_ppt_res.json()
        assert ppt_meta["target_format"] == "pptx"
        assert 0 <= ppt_meta["feasibility_score"] <= 100
        assert "inference" in ppt_meta and len(ppt_meta["inference"]) > 10
        assert "metrics" in ppt_meta
        assert len(ppt_meta["conversion_options"]) >= 2

        feasibility_xls_res = client.get(f"/api/document/{doc_id}/conversion-feasibility?target=excel")
        assert feasibility_xls_res.status_code == 200
        xls_meta = feasibility_xls_res.json()
        assert xls_meta["target_format"] == "excel"
        assert "inference" in xls_meta

        # 6. Direct Upload Feasibility Assessment
        check_res = client.post(
            "/api/check-conversion-feasibility",
            files={"file": ("check.pdf", pdf_bytes, "application/pdf")},
            data={"target": "pptx"}
        )
        assert check_res.status_code == 200
        assert check_res.json()["target_format"] == "pptx"

    def test_structured_pptx_generation(self):
        """Verify structured PowerPoint outputs 16:9 widescreen presentation with native shapes and tables."""
        # Create a document with a title and a table
        doc = fitz.open()
        p1 = doc.new_page(width=792, height=612)
        p1.insert_text((50, 60), "Quarterly Business Review 2026", fontsize=24)
        p1.insert_text((50, 100), "Strategic Initiatives and Accomplishments", fontsize=16)

        p2 = doc.new_page(width=792, height=612)
        p2.insert_text((50, 60), "Key Operational Highlights", fontsize=22)
        p2.insert_text((50, 110), "• Successfully deployed new core infrastructure", fontsize=13)
        p2.insert_text((50, 135), "• Reduced p99 query latency by 45%", fontsize=13)

        # Draw a table grid
        p2.draw_rect(fitz.Rect(50, 200, 450, 300), color=(0, 0, 0), width=1)
        p2.draw_line(fitz.Point(50, 240), fitz.Point(450, 240), color=(0, 0, 0), width=1)
        p2.draw_line(fitz.Point(200, 200), fitz.Point(200, 300), color=(0, 0, 0), width=1)
        p2.insert_text((60, 225), "Category")
        p2.insert_text((210, 225), "Metric")
        p2.insert_text((60, 275), "Uptime")
        p2.insert_text((210, 275), "99.99%")

        pdf_bytes = doc.tobytes()
        doc.close()

        # 1. Check Feasibility
        feasibility = PDFEngine.assess_conversion_feasibility(pdf_bytes, "pptx")
        assert feasibility["feasibility_score"] >= 75
        assert feasibility["feasibility_level"] == "high"
        assert "PowerPoint" in feasibility["inference"]

        # 2. Structured PPTX generation
        pptx_bytes = PDFEngine.convert_pdf_to_pptx(pdf_bytes, mode="structured")
        assert len(pptx_bytes) > 0
        prs = Presentation(io.BytesIO(pptx_bytes))
        assert len(prs.slides) == 2
        # Verify 16:9 widescreen dimensions (13.333 inches = 12192000 EMUs)
        assert round(prs.slide_width.inches, 2) == 13.33
        assert round(prs.slide_height.inches, 1) == 7.5

        # Check slide 1 has title shape
        slide1_text = " ".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
        assert "Quarterly Business Review" in slide1_text

        # Check slide 2 has native table shape
        has_table_shape = any(shape.has_table for shape in prs.slides[1].shapes)
        assert has_table_shape

    def test_excel_conversion_modes(self):
        """Verify multi-sheet and consolidated Excel conversion modes."""
        pdf_bytes = create_sample_pdf(3)

        # 1. Multi-sheet mode (default)
        excel_multi = PDFEngine.convert_pdf_to_excel(pdf_bytes, mode="multi_sheet")
        wb_multi = openpyxl.load_workbook(io.BytesIO(excel_multi))
        assert len(wb_multi.sheetnames) == 3
        assert "Page 1" in wb_multi.sheetnames
        assert "Page 2" in wb_multi.sheetnames
        assert "Page 3" in wb_multi.sheetnames

        # 2. Consolidated mode
        excel_consolidated = PDFEngine.convert_pdf_to_excel(pdf_bytes, mode="consolidated")
        wb_cons = openpyxl.load_workbook(io.BytesIO(excel_consolidated))
        assert len(wb_cons.sheetnames) == 1
        assert "Consolidated" in wb_cons.sheetnames
